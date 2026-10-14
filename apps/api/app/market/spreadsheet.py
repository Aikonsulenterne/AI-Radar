"""Regneark til Til OK: kroner uden OK's tal i radaren.

Radaren kører ikke i et internt miljø og må aldrig se OK's tal (beslutning
2026-10-14). Arket indeholder kun radarens offentlige fund — problem,
teknologi, case, effekten som kilden skriver den og hvad OK skal slå op —
og formler. OK's tal skrives i de gule felter på OK's egen pc; arket
sendes aldrig tilbage.
"""

import re
from collections.abc import Iterable
from dataclasses import dataclass
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.market.catalog import PROBLEMS_BY_KEY

_NUMBER = r"(\d{1,3}(?:[.,]\d+)?)"
_PERCENT = re.compile(_NUMBER + r"\s*(?:%|procent)", re.IGNORECASE)
# "30-40 %" og "30 til 40 procent": den første værdi tæller også.
_RANGE = re.compile(
    _NUMBER + r"\s*%?\s*(?:-|–|til)\s*" + _NUMBER + r"\s*(?:%|procent)", re.IGNORECASE
)

# Andele skrevet med ord, som de står i kilderne.
_WORDS = (
    (re.compile(r"\bhalver(?:et|es|ede|ing)\b", re.IGNORECASE), 50.0),
    (re.compile(r"\bhvert? (?:anden|andet)\b", re.IGNORECASE), 50.0),
    (re.compile(r"\bhvert? (?:tredje)\b", re.IGNORECASE), 33.0),
    (re.compile(r"\bhvert? (?:fjerde)\b", re.IGNORECASE), 25.0),
    (re.compile(r"\bhvert? (?:femte)\b", re.IGNORECASE), 20.0),
)


def effect_percent(text: str | None) -> float | None:
    """Laveste andel i casens effekt (forsigtigt bud), som brøk — fra "30 %",
    "30-40 procent", "halveret" eller "hvert fjerde". None, når kilden ikke
    oplyser en andel; så skriver OK sit eget bud."""
    if not text:
        return None
    found = [m.group(1) for m in _PERCENT.finditer(text)]
    found += [m.group(1) for m in _RANGE.finditer(text)]
    values = [float(v.replace(",", ".")) for v in found]
    values += [value for pattern, value in _WORDS if pattern.search(text)]
    values = [v for v in values if 0 < v <= 100]
    return round(min(values) / 100, 4) if values else None


@dataclass
class SheetRow:
    problem_key: str
    area_label: str
    problem: str
    technology: str
    case: str
    effect: str
    source_url: str | None
    effect_pct: float | None


_INPUT = PatternFill("solid", fgColor="FFF2B3")
_HEAD = PatternFill("solid", fgColor="7A1F35")
_WHITE_BOLD = Font(bold=True, color="FFFFFF")
_BOLD = Font(bold=True)
_WRAP = Alignment(wrap_text=True, vertical="top")

_COLUMNS = (
    ("Område", 14),
    ("Problem", 30),
    ("Teknologi", 24),
    ("Case", 24),
    ("Effekt i casen (som kilden skriver den)", 40),
    ("Kilde", 30),
    ("Slå op internt", 44),
    ("Effekt for OK (ret gerne)", 14),
    ("Hvad tæller I", 26),
    ("Jeres antal pr. år", 14),
    ("Pr. stk.", 30),
    ("Jeres tal pr. stk.", 14),
    ("Kr. pr. år (beregnet)", 18),
)


def _formula(calc: str, row: int) -> str:
    h, j, lcell, rate = f"H{row}", f"J{row}", f"L{row}", "$C$5"
    needed = [h, j, lcell] + ([] if calc == "vaerdi" else [rate])
    blank = ",".join(f'{c}=""' for c in needed)
    product = {
        "tid": f"{j}*{lcell}/60*{rate}*{h}",
        "timer": f"{j}*{lcell}*{rate}*{h}",
        "vaerdi": f"{j}*{lcell}*{h}",
    }[calc]
    return f'=IF(OR({blank}),"",ROUND({product},0))'


def build_workbook(rows: Iterable[SheetRow], month_label: str | None) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Regnestykke"

    ws["A1"] = "Hvad kan det betyde for OK i kroner?"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (
        f"Fra Til OK{f' · {month_label}' if month_label else ''}. "
        "Udfyld de gule felter på OK's egen pc. Radaren ser aldrig jeres tal."
    )
    ws["A3"] = (
        "Effekten kommer fra en anden virksomhed. Jeres resultat kan blive et andet — "
        "brug et forsigtigt bud."
    )
    ws["A5"] = "Jeres timepris (kr. pr. medarbejdertime, inkl. løn og overhead)"
    ws["A5"].font = _BOLD
    ws["C5"].fill = _INPUT
    ws["C5"].number_format = "#,##0"

    head = 7
    for col, (title, width) in enumerate(_COLUMNS, start=1):
        cell = ws.cell(row=head, column=col, value=title)
        cell.fill = _HEAD
        cell.font = _WHITE_BOLD
        cell.alignment = _WRAP
        ws.column_dimensions[get_column_letter(col)].width = width

    r = head
    for item in rows:
        r += 1
        problem = PROBLEMS_BY_KEY.get(item.problem_key)
        calc = problem.calc if problem else "tid"
        values = (
            item.area_label,
            item.problem,
            item.technology,
            item.case,
            item.effect,
            item.source_url or "",
            problem.lookup if problem else "",
            item.effect_pct,
            problem.count_label if problem else "Antal pr. år",
            problem.unit_label if problem else "Pr. stk.",
        )
        for col, value in zip((1, 2, 3, 4, 5, 6, 7, 8, 9, 11), values, strict=True):
            cell = ws.cell(row=r, column=col, value=value)
            cell.alignment = _WRAP
        if item.source_url:
            ws.cell(row=r, column=6).hyperlink = item.source_url
        ws.cell(row=r, column=8).number_format = "0%"
        for col in (8, 10, 12):
            ws.cell(row=r, column=col).fill = _INPUT
        ws.cell(row=r, column=10).number_format = "#,##0"
        ws.cell(row=r, column=12).number_format = "#,##0.0"
        total = ws.cell(row=r, column=13, value=_formula(calc, r))
        total.number_format = "#,##0 kr."
        total.font = _BOLD

    first, last = head + 1, r
    ws.cell(row=r + 2, column=12, value="I alt").font = _BOLD
    sum_cell = ws.cell(
        row=r + 2, column=13, value=f"=SUM(M{first}:M{last})" if last >= first else 0
    )
    sum_cell.number_format = "#,##0 kr."
    sum_cell.font = _BOLD
    ws.freeze_panes = f"C{head + 1}"

    guide = wb.create_sheet("Sådan bruger du arket")
    lines = (
        "Sådan bruger du arket",
        "",
        "1. Skriv jeres timepris i det gule felt øverst (løn og overhead pr. medarbejdertime).",
        "2. Slå tallene i kolonnen »Slå op internt« op i FocalScope eller hos controlling.",
        "3. Skriv dem i de gule felter. Kroner pr. år regnes ud automatisk.",
        "4. »Effekt for OK« er den laveste procent, casen oplyser. Ret den til jeres eget bud.",
        "   Står feltet tomt, oplyser casen ikke en procent — skriv et forsigtigt bud.",
        "",
        "Vigtigt",
        "- Arket bliver på OK's egen pc. Send det ikke tilbage til radaren.",
        "- Radaren gemmer ingen OK-tal og kører ikke i et internt miljø.",
        "- Effekten er målt hos en anden virksomhed. Det er et regneeksempel, ikke en lovning.",
    )
    for i, line in enumerate(lines, start=1):
        line_cell = guide.cell(row=i, column=1, value=line)
        if line in ("Sådan bruger du arket", "Vigtigt"):
            line_cell.font = Font(bold=True, size=13 if i == 1 else 11)
    guide.column_dimensions["A"].width = 100

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
